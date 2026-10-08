"""Excel export for any registered report (openpyxl, in-memory). Rows are capped to keep requests bounded."""

import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.config import settings
from app.services.reports import Col

MAX_EXPORT_ROWS = 50_000

_FORMATS = {
    "qty": "#,##,##0.000",
    "money": "₹ #,##,##0.00",
    "rate": "#,##,##0.0000",
    "pct": '0.00"%"',
    "int": "#,##,##0",
    "date": "dd-mmm-yyyy",
    "datetime": "dd-mmm-yyyy hh:mm",
}


def to_xlsx(title: str, filters: dict, cols: list[Col], rows: list[dict], totals: dict, generated_by: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([f"{settings.COMPANY_NAME} — {title}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generated {datetime.now().strftime('%d-%b-%Y %H:%M')} by {generated_by}"])
    applied = ", ".join(f"{k}: {v}" for k, v in filters.items() if v not in (None, ""))
    ws.append([f"Filters: {applied or 'none'}"])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append([c.label for c in cols])
    fill = PatternFill("solid", fgColor="1F3A5F")
    for i in range(1, len(cols) + 1):
        cell = ws.cell(row=header_row, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r in rows[:MAX_EXPORT_ROWS]:
        values = []
        for c in cols:
            v = r.get(c.key)
            if isinstance(v, Decimal):
                v = float(v)
            if isinstance(v, datetime) and v.tzinfo is not None:
                v = v.replace(tzinfo=None)
            values.append(v)
        ws.append(values)
        row = ws.max_row
        for i, c in enumerate(cols, start=1):
            if c.type in _FORMATS:
                ws.cell(row=row, column=i).number_format = _FORMATS[c.type]

    if totals and rows:
        ws.append([("Total" if i == 0 else (totals.get(c.key))) for i, c in enumerate(cols)])
        row = ws.max_row
        for i, c in enumerate(cols, start=1):
            cell = ws.cell(row=row, column=i)
            cell.font = Font(bold=True)
            if c.type in _FORMATS:
                cell.number_format = _FORMATS[c.type]

    if len(rows) > MAX_EXPORT_ROWS:
        ws.append([f"Truncated to {MAX_EXPORT_ROWS:,} rows. Narrow the filters to export the rest."])

    for i, c in enumerate(cols, start=1):
        sample = [len(str(c.label))] + [len(str(r.get(c.key) or "")) for r in rows[:200]]
        ws.column_dimensions[get_column_letter(i)].width = min(max(sample) + 3, 45)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fmt(v) -> str:
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    return str(v)
